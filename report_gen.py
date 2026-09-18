# report_gen.py — PDF (reportlab) and Excel (openpyxl) report generation
import datetime
from blast_engine import BlastInputs, BlastResults


def generate_pdf_report(inputs: BlastInputs, results: BlastResults, path: str):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table,
            TableStyle, HRFlowable,
        )
        from reportlab.lib.units import cm
    except ImportError:
        raise ImportError("Run: pip install reportlab")

    doc   = SimpleDocTemplate(path, pagesize=A4,
                               topMargin=2*cm, bottomMargin=2*cm,
                               leftMargin=2*cm, rightMargin=2*cm)
    styles = getSampleStyleSheet()
    story  = []

    RED  = colors.HexColor("#e94560")
    DARK = colors.HexColor("#0f3460")
    GREY = colors.HexColor("#f4f4f8")

    title_style = ParagraphStyle("ttl", fontSize=18, fontName="Helvetica-Bold",
                                  textColor=RED, spaceAfter=4)
    sub_style   = ParagraphStyle("sub", fontSize=9,  fontName="Helvetica",
                                  textColor=colors.grey, spaceAfter=10)
    h2_style    = ParagraphStyle("h2",  fontSize=12, fontName="Helvetica-Bold",
                                  textColor=DARK, spaceBefore=14, spaceAfter=4)

    story.append(Paragraph("⚡  BLAST DESIGN REPORT", title_style))
    story.append(Paragraph(
        f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d  %H:%M')}  ·  "
        "Advanced Blast Design Calculator v2.0",
        sub_style
    ))
    story.append(HRFlowable(width="100%", thickness=2, color=RED))
    story.append(Spacer(1, 0.3*cm))

    def table(data, col_w=None):
        t = Table(data, colWidths=col_w or [9*cm, 7*cm])
        t.setStyle(TableStyle([
            ("BACKGROUND",   (0,0), (-1,0), DARK),
            ("TEXTCOLOR",    (0,0), (-1,0), colors.white),
            ("FONTNAME",     (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE",     (0,0), (-1,0), 10),
            ("FONTNAME",     (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE",     (0,1), (-1,-1), 9),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, GREY]),
            ("GRID",         (0,0), (-1,-1), 0.25, colors.HexColor("#cccccc")),
            ("LEFTPADDING",  (0,0), (-1,-1), 8),
            ("TOPPADDING",   (0,0), (-1,-1), 5),
            ("BOTTOMPADDING",(0,0), (-1,-1), 5),
        ]))
        return t

    # Inputs
    story.append(Paragraph("INPUT PARAMETERS", h2_style))
    story.append(table([
        ["Parameter", "Value"],
        ["Bench Height",      f"{inputs.bench_height} m"],
        ["Hole Diameter",     f"{inputs.hole_diameter} mm"],
        ["Drill Angle",       f"{inputs.drill_angle}°"],
        ["Rock Type",         inputs.rock_type],
        ["Rock Density",      f"{inputs.rock_density} t/m³"],
        ["UCS",               f"{inputs.ucs} MPa"],
        ["Rock Factor A",     f"{inputs.rock_factor}"],
        ["Explosive Type",    inputs.explosive_type],
        ["Explosive Density", f"{inputs.explosive_density} g/cm³"],
        ["RWS",               f"{inputs.rws}"],
        ["VOD",               f"{inputs.vod} m/s"],
        ["Burden Factor Kb",  f"{inputs.burden_factor}"],
        ["Spacing Ratio",     f"{inputs.spacing_ratio}"],
        ["Rows × Holes/Row",  f"{inputs.num_rows} × {inputs.holes_per_row}"],
    ]))
    story.append(Spacer(1, 0.4*cm))

    # Results
    story.append(Paragraph("BLAST DESIGN RESULTS", h2_style))
    story.append(table([
        ["Parameter", "Value"],
        ["Burden (B)",             f"{results.burden:.3f} m"],
        ["Spacing (S)",            f"{results.spacing:.3f} m"],
        ["Stemming (T)",           f"{results.stemming:.3f} m"],
        ["Subdrill (J)",           f"{results.subdrill:.3f} m"],
        ["Hole Depth",             f"{results.hole_depth:.3f} m"],
        ["Charge Length",          f"{results.charge_length:.3f} m"],
        ["Langefors Burden (ref)", f"{results.langefors_burden:.3f} m"],
        ["Charge per Hole",        f"{results.charge_per_hole:.2f} kg"],
        ["Total Holes",            f"{results.total_holes}"],
        ["Total Explosives",       f"{results.total_explosives:.1f} kg"],
        ["Total Drill Meters",     f"{results.total_drill_meters:.1f} m"],
        ["Total Volume",           f"{results.total_volume:.1f} m³"],
        ["Total Tonnage",          f"{results.total_tonnage:.1f} t"],
        ["Powder Factor",          f"{results.powder_factor:.4f} kg/m³"],
        ["X₅₀ Fragment Size",      f"{results.x50:.1f} mm"],
        ["X₈₀ Fragment Size",      f"{results.x80:.1f} mm"],
        ["Uniformity Index n",     f"{results.uniformity_index:.3f}"],
        ["Fines < 50 mm",          f"{results.percent_fines:.1f}%"],
        ["Critical PPV Distance",  f"{results.critical_distance:.0f} m"],
        ["Flyrock Distance",       f"{results.flyrock_distance:.0f} m"],
        ["Min Safe Distance",      f"{results.min_safe_distance:.0f} m"],
        ["Total Blast Cost",       f"${results.total_cost:,.2f} USD"],
        ["Cost per Tonne",         f"${results.cost_per_tonne:.3f} USD/t"],
    ]))

    if results.warnings:
        story.append(Spacer(1, 0.4*cm))
        story.append(Paragraph("⚠  DESIGN WARNINGS", h2_style))
        for w in results.warnings:
            story.append(Paragraph(f"• {w}", styles["Normal"]))

    story.append(Spacer(1, 0.5*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.grey))
    story.append(Paragraph(
        "Advanced Blast Design Calculator v2.0 — Mining Engineering Tool",
        sub_style
    ))
    doc.build(story)


def generate_excel_report(inputs: BlastInputs, results: BlastResults, path: str):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise ImportError("Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Blast Design Report"

    RED  = "E94560"
    DARK = "0F3460"
    LITE = "F0F4FF"
    PINK = "FFF0F4"

    def hdr(row, col, text, bg=DARK):
        c = ws.cell(row=row, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=bg)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 20

    def row(r, label, val, unit="", shade=LITE):
        ws.cell(row=r, column=1, value=label).font = Font(size=9)
        v = ws.cell(row=r, column=2, value=round(val, 4) if isinstance(val, float) else val)
        v.font = Font(size=9)
        u = ws.cell(row=r, column=3, value=unit)
        u.font = Font(italic=True, color="888888", size=9)
        if r % 2 == 0:
            for col in range(1, 4):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
        ws.row_dimensions[r].height = 18

    ws.merge_cells("A1:D1")
    t = ws.cell(row=1, column=1,
                value=f"⚡ BLAST DESIGN REPORT — {datetime.datetime.now().strftime('%Y-%m-%d')}")
    t.font = Font(bold=True, size=15, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=RED)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 38

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12

    r = 3
    hdr(r, 1, "INPUT PARAMETERS"); hdr(r, 2, "VALUE"); hdr(r, 3, "UNIT"); r += 1
    inp_rows = [
        ("Bench Height",         inputs.bench_height,        "m"),
        ("Hole Diameter",        inputs.hole_diameter,        "mm"),
        ("Drill Angle",          inputs.drill_angle,          "°"),
        ("Num Rows",             inputs.num_rows,             ""),
        ("Holes per Row",        inputs.holes_per_row,        ""),
        ("Rock Type",            inputs.rock_type,            ""),
        ("Rock Density",         inputs.rock_density,         "t/m³"),
        ("UCS",                  inputs.ucs,                  "MPa"),
        ("Rock Factor A",        inputs.rock_factor,          ""),
        ("Explosive Type",       inputs.explosive_type,       ""),
        ("Explosive Density",    inputs.explosive_density,    "g/cm³"),
        ("RWS",                  inputs.rws,                  ""),
        ("VOD",                  inputs.vod,                  "m/s"),
        ("Burden Factor Kb",     inputs.burden_factor,        ""),
        ("Spacing Ratio S/B",    inputs.spacing_ratio,        ""),
        ("Stemming Factor Kt",   inputs.stemming_factor,      ""),
        ("Subdrill Factor Kj",   inputs.subdrill_factor,      ""),
        ("Explosive Cost",       inputs.explosive_cost_per_kg,"$/kg"),
        ("Drilling Cost",        inputs.drilling_cost_per_m,  "$/m"),
    ]
    for lbl, val, unit in inp_rows:
        row(r, lbl, val, unit, LITE); r += 1

    r += 1
    hdr(r, 1, "CALCULATED RESULTS", RED); hdr(r, 2, "VALUE", RED); hdr(r, 3, "UNIT", RED); r += 1
    res_rows = [
        ("Burden (B)",            results.burden,               "m"),
        ("Spacing (S)",           results.spacing,              "m"),
        ("Stemming (T)",          results.stemming,             "m"),
        ("Subdrill (J)",          results.subdrill,             "m"),
        ("Hole Depth",            results.hole_depth,           "m"),
        ("Charge Length",         results.charge_length,        "m"),
        ("Langefors Burden (ref)",results.langefors_burden,     "m"),
        ("Charge per Hole",       results.charge_per_hole,      "kg"),
        ("Linear Charge Density", results.linear_charge_density,"kg/m"),
        ("Borehole Pressure",     results.borehole_pressure,    "MPa"),
        ("Total Holes",           results.total_holes,          ""),
        ("Total Drill Meters",    results.total_drill_meters,   "m"),
        ("Total Explosives",      results.total_explosives,     "kg"),
        ("Blast Area",            results.blast_area,           "m²"),
        ("Volume per Hole",       results.volume_per_hole,      "m³"),
        ("Total Volume",          results.total_volume,         "m³"),
        ("Total Tonnage",         results.total_tonnage,        "t"),
        ("Powder Factor",         results.powder_factor,        "kg/m³"),
        ("Specific Drilling",     results.specific_drilling,    "m/m³"),
        ("X₅₀ Fragment Size",     results.x50,                  "mm"),
        ("X₈₀ Fragment Size",     results.x80,                  "mm"),
        ("X₂₀ Fragment Size",     results.x20,                  "mm"),
        ("Uniformity Index n",    results.uniformity_index,     ""),
        ("Fines < 50mm",          results.percent_fines,        "%"),
        ("Critical PPV Distance", results.critical_distance,    "m"),
        ("Scaled Distance Limit", results.scaled_distance_limit,""),
        ("Flyrock Distance",      results.flyrock_distance,     "m"),
        ("Min Safe Distance",     results.min_safe_distance,    "m"),
        ("Air Blast Zone",        results.airblast_distance,    "m"),
        ("Drilling Cost",         results.drilling_cost,        "USD"),
        ("Explosive Cost",        results.explosive_cost_total, "USD"),
        ("Initiation Cost",       results.initiation_cost,      "USD"),
        ("Total Blast Cost",      results.total_cost,           "USD"),
        ("Cost per Tonne",        results.cost_per_tonne,       "USD/t"),
        ("Cost per m³",           results.cost_per_m3,          "USD/m³"),
    ]
    for lbl, val, unit in res_rows:
        row(r, lbl, val, unit, PINK); r += 1

    if results.warnings:
        r += 1
        hdr(r, 1, "DESIGN WARNINGS", "CC3300"); r += 1
        for w in results.warnings:
            ws.cell(row=r, column=1, value=f"⚠ {w}").font = Font(color="CC3300"); r += 1

    wb.save(path)